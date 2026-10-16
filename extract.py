import sys,json,subprocess,zipfile,xml.etree.ElementTree as ET
from pathlib import Path
CAP=300000

def extract(path):
    p=Path(path); e=p.suffix.lower(); text=''; meta={}; state='indexed'
    if e=='.pdf':
        # Limit pages and stream output into a bounded buffer; no shell, no macros.
        proc=subprocess.Popen(['pdftotext','-f','1','-l','500','-layout',str(p),'-'],stdout=subprocess.PIPE,stderr=subprocess.DEVNULL)
        raw=proc.stdout.read(CAP*4); proc.stdout.close(); proc.terminate(); proc.wait(timeout=3)
        text=raw.decode('utf-8','replace'); meta['extraction_limit']='first 500 pages, 300000 characters'
        if not text.strip(): state='needs_ocr'
    elif e=='.docx':
        from docx import Document
        d=Document(p); chunks=[]
        for para in d.paragraphs: chunks.append(para.text)
        for table in d.tables:
            for row in table.rows: chunks.append(' | '.join(c.text for c in row.cells))
        text='\n'.join(chunks)
    elif e in ('.xlsx','.xlsm'):
        from openpyxl import load_workbook
        wb=load_workbook(p,read_only=True,data_only=True,keep_links=False); parts=[]; size=0
        for ws in wb:
            parts.append('\nSheet: '+ws.title)
            for row in ws.iter_rows():
                vals=['%s=%s'%(c.coordinate,str(c.value)) for c in row if c.value is not None]
                line=' | '.join(vals); parts.append(line); size+=len(line)
                if size>=CAP: break
            if size>=CAP: break
        wb.close(); text='\n'.join(parts); meta['spreadsheet_values']='stored/cached values; formulas not recalculated'
    elif e=='.xls':
        import xlrd
        wb=xlrd.open_workbook(str(p),on_demand=True); parts=[]; size=0
        for ws in wb.sheets():
            parts.append('\nSheet: '+ws.name)
            for i in range(ws.nrows):
                line='Row %s: '%(i+1)+' | '.join('Col %s=%s'%(j+1,v) for j,v in enumerate(ws.row_values(i)) if v!=''); parts.append(line); size+=len(line)
                if size>=CAP: break
            if size>=CAP: break
        wb.release_resources(); text='\n'.join(parts)
    elif e=='.pptx':
        from pptx import Presentation
        prs=Presentation(str(p)); text='\n'.join('Slide %s: '%(i+1)+'\n'.join(s.text for s in slide.shapes if s.has_text_frame) for i,slide in enumerate(prs.slides))
    elif e in ('.doc','.ppt','.rtf'):
        cmd={'.doc':['antiword'],'.ppt':['catppt'],'.rtf':['unrtf','--text']}[e]
        proc=subprocess.Popen(cmd+[str(p)],stdout=subprocess.PIPE,stderr=subprocess.DEVNULL)
        text=proc.stdout.read(CAP*4).decode('utf-8','replace'); proc.stdout.close(); proc.terminate(); proc.wait(timeout=3)
    elif e in ('.odt','.ods'):
        with zipfile.ZipFile(p) as z:
            if z.getinfo('content.xml').file_size>30_000_000: raise ValueError('XML member too large')
            text=' '.join(ET.fromstring(z.read('content.xml')).itertext())
    elif e in ('.txt','.md','.csv','.html','.htm','.sql'):
        with p.open('rb') as f: text=f.read(CAP*4).decode('utf-8','replace')
    elif e in {'.mp3','.flac','.m4a','.ogg','.wav','.wma','.aac','.opus'}:
        from mutagen import File
        audio=File(p,easy=True)
        if audio:
            meta={str(k):v for k,v in (audio.tags or {}).items()}; meta['duration_seconds']=round(audio.info.length,2)
    elif e in {'.jpg','.jpeg','.png','.heic','.tif','.tiff','.gif','.webp','.cr2','.dng'}:
        r=subprocess.run(['exiftool','-json','-DateTimeOriginal','-CreateDate','-ImageWidth','-ImageHeight','-Make','-Model','-Description','-Subject','-Keywords','-GPSLatitude','-GPSLongitude',str(p)],capture_output=True,timeout=12)
        meta=json.loads(r.stdout or b'[{}]')[0]; meta.pop('SourceFile',None)
    elif e in {'.mp4','.mkv','.avi','.mov','.m4v','.mpg','.mpeg','.webm','.3gp'}:
        r=subprocess.run(['ffprobe','-v','quiet','-show_format','-of','json',str(p)],capture_output=True,timeout=12)
        meta=json.loads(r.stdout or b'{}').get('format',{}); meta.pop('filename',None)
    else: state='metadata_only'
    meta['text_truncated']=len(text)>CAP
    return {'text':text[:CAP],'metadata':json.dumps(meta,ensure_ascii=False),'state':state}

if __name__=='__main__':
    try: print(json.dumps(extract(sys.argv[1]),ensure_ascii=True))
    except Exception as e: print(json.dumps({'text':'','metadata':'','state':'failed','error':type(e).__name__+': '+str(e)[:250]}))
